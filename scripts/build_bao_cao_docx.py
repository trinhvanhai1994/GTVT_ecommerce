#!/usr/bin/env python3
"""Build BAO-CAO-DO-AN.docx from BAO-CAO-DO-AN.md (Mermaid PNG + Pandoc + post-process)."""

from __future__ import annotations

import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC_MD = ROOT / "BAO-CAO-DO-AN.md"
EXPORT_DIR = ROOT / "build" / "docx-export"
DIAGRAM_DIR = EXPORT_DIR / "diagrams"
EXPORT_MD = EXPORT_DIR / "BAO-CAO-DO-AN_export.md"
REFERENCE_DOC = Path(__file__).resolve().parent / "reference_bao_cao.docx"
PUPPETEER_CONFIG = Path(__file__).resolve().parent / "puppeteer-config.json"
OUT_DOCX = ROOT / "BAO-CAO-DO-AN.docx"
PANDOC = "/opt/homebrew/bin/pandoc"

MERMAID_RE = re.compile(r"```mermaid\s*\n(.*?)```", re.DOTALL)
CHAPTER_H1 = re.compile(r"^[1-9]\.\s")
PHU_LUC = re.compile(r"Phụ\s+lục", re.IGNORECASE)
SECTION_H2 = re.compile(r"^\d+\.\d+\s")
TABLE_SEP = re.compile(r"^\|[-:\s|]+\|$")
COVER_MARKERS = ("BÁO CÁO ĐỒ ÁN", "Tên đề tài", "Xây dựng hệ thống")

HINH_BOLD_MERMAID = re.compile(
    r"(\*\*Hình\s+([\d.a-z]+)\s+—[^\n]*\*\*\s*\n)\s*```mermaid\s*\n(.*?)```",
    re.DOTALL | re.IGNORECASE,
)
HINH_HEADING_MERMAID = re.compile(
    r"(^#{1,3}\s+[^\n]*Hình\s+([\d.a-z]+)\s+—[^\n]*\s*\n)\s*```mermaid\s*\n(.*?)```",
    re.MULTILINE | re.DOTALL | re.IGNORECASE,
)


def run(cmd: list[str], cwd: Path | None = None, env: dict | None = None) -> None:
    print("+", " ".join(cmd))
    merged = os.environ.copy()
    merged.setdefault("PUPPETEER_CACHE_DIR", str(ROOT / "build" / "puppeteer-cache"))
    if env:
        merged.update(env)
    subprocess.run(cmd, cwd=cwd, check=True, env=merged)


def count_md_pipe_tables(text: str) -> int:
    lines = text.splitlines()
    count = 0
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith("|") and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1].strip()):
            count += 1
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                i += 1
            continue
        i += 1
    return count


def hinh_to_filename(label: str) -> str:
    return "hinh-" + label.strip().lower().replace(".", "-")


def mermaid_scale_for_id(diagram_id: str) -> list[str] | None:
    if diagram_id in ("hinh-7-1", "hinh-6-6", "hinh-6-1"):
        return ["-s", "3"]
    return None


def render_mermaid(diagram_id: str, source: str, extra_args: list[str] | None = None) -> Path:
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    mmd = DIAGRAM_DIR / f"{diagram_id}.mmd"
    png = DIAGRAM_DIR / f"{diagram_id}.png"
    mmd.write_text(source.strip() + "\n", encoding="utf-8")
    if png.is_file() and png.stat().st_mtime >= mmd.stat().st_mtime:
        return png
    args = [
        "npx",
        "-y",
        "@mermaid-js/mermaid-cli",
        "-i",
        str(mmd),
        "-o",
        str(png),
        "-b",
        "white",
    ]
    if PUPPETEER_CONFIG.is_file():
        args.extend(["-p", str(PUPPETEER_CONFIG)])
    if extra_args:
        args.extend(extra_args)
    else:
        args.extend(["-s", "2"])
    run(args)
    return png


def demote_front_matter_headings(text: str) -> str:
    out: list[str] = []
    for line in text.splitlines():
        if re.match(r"^#\s+BÁO CÁO", line, re.I):
            out.append(f"**{line.lstrip('# ').strip()}**")
        elif re.match(r"^##\s+Mục lục\s*$", line, re.I):
            out.append("**Mục lục**")
        else:
            out.append(line)
    return "\n".join(out) + "\n"


def replace_mermaid_with_png(text: str) -> str:
    def repl_bold(m: re.Match[str]) -> str:
        label = m.group(2)
        diagram_id = hinh_to_filename(label)
        render_mermaid(diagram_id, m.group(3), mermaid_scale_for_id(diagram_id))
        return f"{m.group(1)}\n\n![](diagrams/{diagram_id}.png)\n"

    def repl_head(m: re.Match[str]) -> str:
        label = m.group(2)
        diagram_id = hinh_to_filename(label)
        render_mermaid(diagram_id, m.group(3), mermaid_scale_for_id(diagram_id))
        return f"{m.group(1)}\n\n![](diagrams/{diagram_id}.png)\n"

    text = HINH_BOLD_MERMAID.sub(repl_bold, text)
    text = HINH_HEADING_MERMAID.sub(repl_head, text)

    idx = 0
    for m in list(MERMAID_RE.finditer(text)):
        idx += 1
        diagram_id = f"diagram-{idx:02d}"
        render_mermaid(diagram_id, m.group(1), mermaid_scale_for_id(diagram_id))
        text = text.replace(m.group(0), f"\n\n![](diagrams/{diagram_id}.png)\n", 1)

    return text


def prepare_export_md(text: str) -> str:
    text = demote_front_matter_headings(text)
    text = re.sub(
        r"(\*\*Mục lục\*\*\s*\n)\s*(?:\d+\.\s*\[[^\]]+\]\([^\)]+\)\s*\n)+",
        r"\1\n",
        text,
        count=1,
        flags=re.IGNORECASE,
    )
    text = replace_mermaid_with_png(text)
    return text


def _set_run_font(run, size: Pt) -> None:
    run.font.name = "Times New Roman"
    run.font.size = size
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    existing = tbl_pr.find(qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tbl_pr.append(borders)


def table_font_size(table) -> Pt:
    ncols = len(table.columns)
    if ncols >= 4:
        return Pt(11)
    if ncols == 3 and table.rows:
        h0 = table.rows[0].cells[0].text.strip()
        if h0 in ("Mục", "FR", "ID", "Thuật ngữ", "Tầng", "FR (mục 3.2)"):
            return Pt(11)
    return Pt(12)


def row_is_empty(row) -> bool:
    return all(not cell.text.strip() for cell in row.cells)


def remove_trailing_empty_rows(table) -> None:
    while len(table.rows) > 1 and row_is_empty(table.rows[-1]):
        table._tbl.remove(table.rows[-1]._tr)


def usable_text_width(doc: Document) -> int:
    section = doc.sections[0]
    return int(section.page_width - section.left_margin - section.right_margin)


def format_tables(doc: Document) -> None:
    width = usable_text_width(doc)
    for table in doc.tables:
        remove_trailing_empty_rows(table)
        try:
            table.style = "Table Grid"
        except KeyError:
            pass
        set_table_borders(table)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.width = width
        size = table_font_size(table)
        for ri, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                for p in cell.paragraphs:
                    for run in p.runs:
                        _set_run_font(run, size)
                        if ri == 0:
                            run.bold = True


def ensure_reference_doc(force: bool = False) -> None:
    if REFERENCE_DOC.exists() and not force:
        return
    if force and REFERENCE_DOC.exists():
        REFERENCE_DOC.unlink()

    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    for style_name, size in [("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 13)]:
        st = doc.styles[style_name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        st.paragraph_format.line_spacing = 1.15

    sample = doc.add_table(rows=2, cols=2)
    sample.style = "Table Grid"
    set_table_borders(sample)
    REFERENCE_DOC.parent.mkdir(parents=True, exist_ok=True)
    doc.save(REFERENCE_DOC)
    print(f"Wrote {REFERENCE_DOC}")


def _style_name(paragraph) -> str:
    return (paragraph.style.name if paragraph.style else "").replace(" ", "").lower()


def is_heading(paragraph, level: int) -> bool:
    return _style_name(paragraph) == f"heading{level}"


def get_paragraph_bookmark(paragraph) -> str | None:
    for bm in paragraph._p.findall(qn("w:bookmarkStart")):
        name = bm.get(qn("w:name"))
        if name and not name.startswith("_GoBack"):
            return name
    return None


def _max_bookmark_id(doc: Document) -> int:
    max_id = 0
    for bm in doc.element.body.iter(qn("w:bookmarkStart")):
        try:
            max_id = max(max_id, int(bm.get(qn("w:id"), "0")))
        except ValueError:
            pass
    return max_id


def ensure_heading_bookmarks(doc: Document) -> None:
    next_id = _max_bookmark_id(doc) + 1
    used_names: set[str] = set()

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        toc_heading = (is_heading(p, 1) and (CHAPTER_H1.search(text) or PHU_LUC.search(text))) or (
            is_heading(p, 2) and SECTION_H2.match(text)
        )
        if not toc_heading or get_paragraph_bookmark(p):
            continue

        slug = re.sub(r"[^\w]+", "_", text.lower()).strip("_")[:48] or "entry"
        name = f"toc_{slug}"
        suffix = 0
        while name in used_names:
            suffix += 1
            name = f"toc_{slug}_{suffix}"
        used_names.add(name)

        bm_start = OxmlElement("w:bookmarkStart")
        bm_start.set(qn("w:id"), str(next_id))
        bm_start.set(qn("w:name"), name)
        bm_end = OxmlElement("w:bookmarkEnd")
        bm_end.set(qn("w:id"), str(next_id))
        p._p.insert(0, bm_start)
        p._p.append(bm_end)
        next_id += 1


def configure_toc_line(paragraph, level: int, tab_pos_twips: int) -> None:
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    if level == 2:
        pf.left_indent = Cm(0.75)
    p_pr = paragraph._p.get_or_add_pPr()
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        p_pr.append(tabs)
    else:
        for t in list(tabs):
            tabs.remove(t)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), str(tab_pos_twips))
    tabs.append(tab)


def append_pageref_field(paragraph, bookmark: str) -> None:
    tab_run = paragraph.add_run()
    tab_run._r.append(OxmlElement("w:tab"))

    r_begin = paragraph.add_run()
    fc_begin = OxmlElement("w:fldChar")
    fc_begin.set(qn("w:fldCharType"), "begin")
    r_begin._r.append(fc_begin)

    r_instr = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f"PAGEREF {bookmark} \\h"
    r_instr._r.append(instr)

    r_sep = paragraph.add_run()
    fc_sep = OxmlElement("w:fldChar")
    fc_sep.set(qn("w:fldCharType"), "separate")
    r_sep._r.append(fc_sep)

    r_text = paragraph.add_run("0")
    _set_run_font(r_text, Pt(13))

    r_end = paragraph.add_run()
    fc_end = OxmlElement("w:fldChar")
    fc_end.set(qn("w:fldCharType"), "end")
    r_end._r.append(fc_end)


def collect_toc_entries(doc: Document) -> list[tuple[int, str, str | None]]:
    entries: list[tuple[int, str, str | None]] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        bm = get_paragraph_bookmark(p)
        if is_heading(p, 1) and (CHAPTER_H1.search(text) or PHU_LUC.search(text)):
            entries.append((1, text, bm))
        elif is_heading(p, 2) and SECTION_H2.match(text):
            entries.append((2, text, bm))
    return entries


def tab_stop_twips(doc: Document) -> int:
    section = doc.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    if hasattr(width, "twips"):
        return int(width.twips)
    return int(int(width) * 1440 / 914400)


def insert_static_toc(doc: Document, toc_title_idx: int, entries: list[tuple[int, str, str | None]]) -> None:
    tab_pos = tab_stop_twips(doc)
    title_p = doc.paragraphs[toc_title_idx]
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title_p.runs:
        run.bold = True
        _set_run_font(run, Pt(14))

    to_remove: list = []
    for j in range(toc_title_idx + 1, len(doc.paragraphs)):
        p = doc.paragraphs[j]
        pt = p.text.strip()
        if not pt:
            to_remove.append(p._element)
            continue
        if is_heading(p, 1) and CHAPTER_H1.search(pt):
            break
        to_remove.append(p._element)

    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)

    anchor = title_p._element
    missing_bm = 0
    for level, title, bookmark in entries:
        p = doc.add_paragraph()
        anchor.addnext(p._element)
        anchor = p._element
        configure_toc_line(p, level, tab_pos)
        run = p.add_run(title)
        _set_run_font(run, Pt(13))
        run.bold = level == 1
        if bookmark:
            append_pageref_field(p, bookmark)
        else:
            missing_bm += 1

    if missing_bm:
        print(f"Warning: {missing_bm} TOC entries without bookmark.")


def add_page_break_after_element(element) -> None:
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    element.addnext(p)


def add_page_break_before(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    pb = OxmlElement("w:pageBreakBefore")
    p_pr.append(pb)


def format_cover_and_headings(doc: Document) -> None:
    for p in doc.paragraphs:
        text = p.text.strip()
        if any(text.startswith(m) for m in COVER_MARKERS):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                _set_run_font(run, Pt(13))

        if is_heading(p, 1) or is_heading(p, 2) or is_heading(p, 3):
            pf = p.paragraph_format
            pf.space_before = Pt(12)
            pf.space_after = Pt(6)
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pf.line_spacing = 1.15
            size = Pt(14) if is_heading(p, 1) else Pt(13)
            for run in p.runs:
                _set_run_font(run, size)
                run.bold = True


def apply_page_breaks(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:tbl"):
            add_page_break_after_element(child)
            break

    for p in doc.paragraphs:
        if not is_heading(p, 1):
            continue
        text = p.text.strip()
        m = re.match(r"^([1-9])\.\s", text)
        if m and int(m.group(1)) >= 2:
            add_page_break_before(p)
        elif PHU_LUC.search(text):
            add_page_break_before(p)


def postprocess_docx(path: Path) -> None:
    doc = Document(str(path))

    for section in doc.sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)

    ensure_heading_bookmarks(doc)
    format_cover_and_headings(doc)
    apply_page_breaks(doc)

    toc_idx: int | None = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().lower() == "mục lục":
            toc_idx = i
            break

    toc_entries = collect_toc_entries(doc)
    if toc_idx is not None and toc_entries:
        insert_static_toc(doc, toc_idx, toc_entries)
        print(f"Static TOC: {len(toc_entries)} entries")

    for p in doc.paragraphs:
        style_name = p.style.name if p.style else ""
        text = p.text.strip()
        if style_name in ("Normal", "Body Text", "First Paragraph"):
            if text.lower() == "mục lục":
                continue
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pf.line_spacing = 1.35
            for run in p.runs:
                if run.font.size is None:
                    _set_run_font(run, Pt(13))
                elif not run.font.name:
                    _set_run_font(run, run.font.size)

        if re.match(r"^Hình\s+[\d.a-z]+", text, re.I):
            for run in p.runs:
                run.bold = True

    for inline in doc.inline_shapes:
        if inline.width > Cm(16):
            ratio = float(Cm(16)) / float(inline.width)
            inline.width = int(inline.width * ratio)
            inline.height = int(inline.height * ratio)

    format_tables(doc)
    doc.save(str(path))
    print(f"Post-processed {path}")


def validate_docx(path: Path, expected_tables: int, expected_mermaid: int = 16) -> None:
    doc = Document(str(path))
    n_tables = len(doc.tables)
    if n_tables != expected_tables:
        raise SystemExit(f"Table count mismatch: MD {expected_tables}, DOCX {n_tables}")

    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    if "flowchart" in xml.lower():
        raise SystemExit("DOCX still contains raw mermaid flowchart text")

    borders = xml.count("w:tblBorders")
    if borders < n_tables:
        raise SystemExit(f"Missing tblBorders: {borders} for {n_tables} tables")

    print(
        f"Validated: {n_tables} tables, {len(media)} images (expected ~{expected_mermaid}), "
        f"PAGEREF {xml.count('PAGEREF')}"
    )


def main() -> int:
    if not SRC_MD.is_file():
        print(f"Missing {SRC_MD}", file=sys.stderr)
        return 1

    raw = SRC_MD.read_text(encoding="utf-8")
    expected_tables = count_md_pipe_tables(raw)
    expected_mermaid = len(re.findall(r"```mermaid", raw))

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    export_text = prepare_export_md(raw)
    EXPORT_MD.write_text(export_text, encoding="utf-8")
    print(f"Wrote {EXPORT_MD} ({expected_tables} tables, {expected_mermaid} mermaid in source)")

    ensure_reference_doc(force=True)

    pandoc_bin = PANDOC if Path(PANDOC).is_file() else "pandoc"
    run(
        [
            pandoc_bin,
            str(EXPORT_MD.name),
            "--from",
            "markdown+pipe_tables+table_captions",
            "--to",
            "docx",
            "--resource-path=.",
            f"--reference-doc={REFERENCE_DOC}",
            "-o",
            str(OUT_DOCX),
        ],
        cwd=EXPORT_DIR,
    )

    postprocess_docx(OUT_DOCX)
    validate_docx(OUT_DOCX, expected_tables, expected_mermaid)
    print(f"Done: {OUT_DOCX}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
